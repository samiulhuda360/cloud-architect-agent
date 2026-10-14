"""Company policies: reading documents, extracting rules, enforcing accepted rules; requirements briefs; the model cache."""

import io
import json
from pathlib import Path

import pytest

from cloud_architect import llm
from cloud_architect.assess import assess
from cloud_architect.brief import extract as extract_brief
from cloud_architect.config import Settings
from cloud_architect.documents import grounded, read_document, sentences
from cloud_architect.evaluate import load_labels, score_extraction
from cloud_architect.models import Component, Design, PolicyRule, Workload
from cloud_architect.policies import extract, extract_with_model, extract_with_rules, places_to_regions, service_ids
from cloud_architect.regions import RESIDENCY_REGIONS
from cloud_architect.rules import architect
from cloud_architect.terraform import generate

ROOT = Path(__file__).resolve().parents[1]
CFG = Settings(llm_api_key="test", llm_cache=None, min_interval=0)


def make_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with plain text, enough for a text extractor to read."""
    esc = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
    stream = "BT /F1 11 Tf 14 TL 50 790 Td " + " ".join(f"({t}) Tj T*" for t in esc) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    import docx

    d = docx.Document()
    for style, text in paragraphs:
        d.add_heading(text, level=1) if style == "h" else d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- reading documents
def test_pdf_and_word_documents_are_read_with_their_sections():
    pdf = read_document("policy.pdf", make_pdf(["4.2 Backups", "Backups must be retained for at least 30 days."]))
    assert "retained for at least 30 days" in pdf
    word = read_document("policy.docx", make_docx([("h", "Hosting"), ("p", "Customer data must be stored only in the United Kingdom.")]))
    s = sentences(word)
    assert s[0].section == "Hosting" and "United Kingdom" in s[0].text
    with pytest.raises(ValueError):
        read_document("policy.exe", b"x")


def test_the_example_pdf_gives_the_same_rules_as_its_markdown():
    pdf = (ROOT / "examples/policies/hollin-wren-security-standard.pdf").read_bytes()
    md = (ROOT / "examples/policies/hollin-wren-security-standard.md").read_bytes()
    checks = [[r["check"] for r in extract(name, data)["rules"]] for name, data in (("s.pdf", pdf), ("s.md", md))]
    assert checks[0] == checks[1] and len(checks[0]) == 7


def test_a_list_after_a_colon_belongs_to_its_sentence():
    s = sentences(
        "## 3. Approved services\n\nTeams may only use the following services:\n\n- App Service\n- Key Vault\n\nNew services need approval."
    )
    assert s[0].section == "3" and s[0].follows == ("App Service", "Key Vault") and s[1].text == "New services need approval."


def test_places_and_services_are_recognised():
    assert set(places_to_regions("only in the United Kingdom")) == {"uksouth", "ukwest"}
    assert places_to_regions("UK South and UK West only") == ["uksouth", "ukwest"]
    assert places_to_regions("tell us in writing") == []  # "us" is not the United States
    assert set(places_to_regions("hosted within the European Union")) == RESIDENCY_REGIONS["eu"]
    assert service_ids("Azure Database for PostgreSQL, Azure Key Vault and Azure Monitor") == ["postgres", "key_vault", "log_analytics"]


# ---------------------------------------------------------------- extraction
def test_rule_based_extraction_finds_every_labelled_rule_in_the_example_policies():
    labels = {d: r for d, r in load_labels().items() if d.startswith("examples/")}
    predicted = {d: [r.model_dump() for r in extract_with_rules((ROOT / d).read_text(encoding="utf-8"), Path(d).name)] for d in labels}
    score = score_extraction(predicted, labels)
    assert score["recall"] == 1.0 and score["precision"] == 1.0, score


def test_every_extracted_rule_quotes_its_document():
    text = (ROOT / "examples/policies/lumenvale-hipaa-security-policy.md").read_text(encoding="utf-8")
    for r in extract_with_rules(text, "lumenvale.md"):
        assert r.status == "proposed" and r.extracted_by == "rules" and grounded(r.quote, text) and r.section


def test_model_rules_whose_quote_is_not_in_the_document_are_dropped(fake_llm):
    text = "## 3. Hosting\n\nCustomer data must be stored only in the United Kingdom.\n"
    reply = {
        "rules": [
            {
                "check": "allowed_regions",
                "places": ["United Kingdom"],
                "section": "3",
                "quote": "Customer data must be stored only in the United Kingdom.",
            },
            {"check": "customer_managed_keys", "section": "5", "quote": "All data uses customer-managed keys."},  # invented
            {"check": "password_length", "quote": "Customer data must be stored only in the United Kingdom."},  # not a check type
        ]
    }
    rules, dropped = extract_with_model(text, "standard.md", CFG, client=fake_llm([json.dumps(reply)]))
    assert [r.check for r in rules] == ["allowed_regions"] and dropped == 1
    assert set(rules[0].regions) == {"uksouth", "ukwest"} and rules[0].extracted_by == "model"


def test_extraction_falls_back_to_rules_when_the_model_fails(fake_llm):
    class Broken:
        chat = type("C", (), {"completions": type("X", (), {"create": staticmethod(lambda **k: (_ for _ in ()).throw(RuntimeError("down")))})()})()

    data = (ROOT / "examples/policies/kestrel-moore-client-data-policy.md").read_bytes()
    out = extract("kestrel.md", data, CFG, client=Broken())
    assert out["method"] == "rules" and "unavailable" in out["note"] and len(out["rules"]) == 4


def test_brief_extraction_prefills_the_form():
    data = (ROOT / "examples/briefs/falkenhain-supplier-portal-rfp.md").read_bytes()
    out = extract_brief("rfp.md", data)
    fields = {f["field"]: f["value"] for f in out["fields"]}
    assert out["method"] == "rules"
    assert fields["monthly_users"] == 1200 and fields["peak_concurrent_users"] == 150 and fields["data_growth_pct"] == 30
    assert fields["rpo_minutes"] == 15 and fields["rto_minutes"] == 240 and fields["availability"] == "high"
    assert set(fields["frameworks"]) == {"gdpr", "iso27001"} and fields["budget_month"] == 3500 and fields["blob_gb"] == 1536
    assert all(f["quote"] for f in out["fields"])
    Workload(name="Falkenhain", **fields)  # every proposed value is one the form accepts


# ---------------------------------------------------------------- enforcement
def rule(check, **kw):
    return PolicyRule(check=check, status="accepted", source="policy.md", section="4", quote=f"The {check} rule.", id="P1", **kw)


def pol(*rules, **kw):
    return Workload(name="Shop", needs=["web_app", "relational_db", "object_storage"], residency="uk", policies=list(rules), **kw)


def violations(d, w, book):
    return {v.rule: v for v in assess(d, w, book).violations}


def base_design():
    return Design(
        region="uksouth",
        components=[
            Component(service="app_service", tier="S1", region="uksouth"),
            Component(service="postgres", tier="D2ds_v5", region="uksouth"),
            Component(service="blob_storage", tier="hot_zrs", region="uksouth"),
            Component(service="redis", tier="standard_c0", region="uksouth"),
            Component(service="log_analytics", tier="pay_as_you_go", region="uksouth"),
        ],
    )


@pytest.mark.parametrize(
    "policy,code",
    [
        (rule("allowed_regions", regions=["ukwest"]), "POL-REGION"),
        (rule("customer_managed_keys"), "POL-CMK"),
        (rule("no_public_endpoints"), "POL-PUBLIC"),
        (rule("min_backup_retention_days", days=60), "POL-BACKUP"),
        (rule("approved_services", services=["app_service", "postgres", "blob_storage", "log_analytics"]), "POL-SERVICE"),
        (rule("max_monthly_spend", amount=100, currency="EUR"), "POL-SPEND"),
        (rule("required_tags", tags=["owner"]), "POL-TAGS"),
        (rule("min_log_retention_days", days=365), "POL-LOGS"),
    ],
)
def test_each_accepted_policy_rule_is_enforced_and_cited(book, policy, code):
    v = violations(base_design(), pol(policy), book)
    assert code in v and v[code].source == 'policy.md, section 4: "' + policy.quote + '"'
    proposed = policy.model_copy(update={"status": "proposed"})
    assert code not in violations(base_design(), pol(proposed), book)  # only accepted rules are enforced


def test_the_rule_based_architect_meets_every_policy_rule(book):
    rules = [
        rule("customer_managed_keys"),
        rule("no_public_endpoints"),
        rule("min_backup_retention_days", days=60),
        rule("approved_services", services=["app_service", "postgres", "blob_storage", "key_vault", "log_analytics", "backup_vault"]),
        rule("required_tags", tags=["owner", "cost-centre"]),
        rule("min_log_retention_days", days=800),
        rule("max_monthly_spend", amount=5000, currency="GBP"),
    ]
    w = pol(*rules, peak_rps=120, read_ratio=0.9)  # read-heavy: a cache would help, but Redis is not approved
    d = architect(w, book)
    a = assess(d, w, book)
    assert not a.violations, [v.message for v in a.violations]
    services = {c.service for c in d.components}
    assert "redis" not in services and "backup_vault" in services
    tf = generate(d, w)
    needles = (
        "azurerm_key_vault_key",
        "customer_managed_key",
        "azurerm_storage_account_customer_managed_key",
        'variable "required_tags"',
        "total_retention_in_days = 1095",
        "azurerm_data_protection_backup_instance_postgresql_flexible_server",
        "tags = local.tags",
    )
    for needle in needles:
        assert needle in tf, needle


def test_spend_caps_are_converted_between_currencies(book):
    d = base_design()
    gbp = assess(d, pol(), book).total_month
    cap_eur = gbp / book.fx("GBP") * book.fx("EUR")
    assert "POL-SPEND" in violations(d, pol(rule("max_monthly_spend", amount=cap_eur * 0.99, currency="EUR")), book)
    assert "POL-SPEND" not in violations(d, pol(rule("max_monthly_spend", amount=cap_eur * 1.01, currency="EUR")), book)


def test_an_exception_can_relax_a_policy_rule_openly(book):
    d = base_design()
    d.accepted_exceptions = ["[POL-TAGS] tags are applied by the platform team's pipeline"]
    a = assess(d, pol(rule("required_tags", tags=["owner"])), book)
    assert "POL-TAGS" not in {v.rule for v in a.violations} and any(f.rule == "POL-TAGS" and "Accepted:" in f.message for f in a.findings)


# ---------------------------------------------------------------- the model cache
def test_model_replies_are_cached_on_disk(tmp_path):
    from openai.types.chat import ChatCompletion

    calls = []
    reply = {
        "id": "r1", "object": "chat.completion", "created": 0, "model": "m",
        "choices": [{"index": 0, "finish_reason": "stop", "message": {"role": "assistant", "content": "hello"}}],
    }  # fmt: skip

    class Client:
        class chat:  # noqa: N801
            class completions:  # noqa: N801
                @staticmethod
                def create(**kw):
                    calls.append(kw)
                    return ChatCompletion.model_validate(reply)

    cache = llm.Cache(tmp_path / "llm.sqlite")
    for _ in range(3):
        resp = llm.complete(Client, cache=cache, model="m", messages=[{"role": "user", "content": "hi"}])
        assert resp.choices[0].message.content == "hello"
    assert len(calls) == 1


def test_a_daily_quota_stops_instead_of_retrying():
    import httpx
    import openai

    def quota(**kw):
        raise openai.RateLimitError(
            "Quota exceeded for metric GenerateRequestsPerDayPerProjectPerModel",
            response=httpx.Response(429, request=httpx.Request("POST", "http://x")),
            body=None,
        )

    client = type("C", (), {"chat": type("X", (), {"completions": type("Y", (), {"create": staticmethod(quota)})()})()})()
    with pytest.raises(llm.DailyQuotaError):
        llm.complete(client, retry_waits=(0, 0), model="m", messages=[])
